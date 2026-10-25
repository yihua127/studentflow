"""In-memory Word exports preserve reviewed content and contain no active resources."""

from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

from studentflow.meeting_docx import meeting_minutes_docx

FIRST = "a" * 64
SECOND = "b" * 64
UNUSED = "c" * 64
QUOTE = "Alex will prepare the project outline."


def sources():
    return {
        FIRST: {"provider": "google_chat", "title": "Project discussion", "sender_name": "Alex",
                "sender_id": "users/private-123", "sent_at": "2026-10-26T01:00:00+00:00",
                "account": "private-account@example.invalid",
                "text": QUOTE + "\nUNCITED PRIVATE SOURCE BODY"},
        SECOND: {"provider": "attachment", "title": "Project brief.pdf",
                 "text": "The report must include the survey findings.\nPRIVATE ATTACHMENT TAIL"},
        UNUSED: {"provider": "google_chat", "title": "UNCITED SOURCE TITLE",
                 "text": "UNUSED SOURCE BODY"},
    }


def summary():
    return {
        "title": "Campus survey project discussion",
        "summary": f"The outline was assigned and the survey findings are required ({FIRST}).",
        "citations": [{"source_id": FIRST, "quote": QUOTE}],
        "decisions": [{"text": "Include the survey findings in the report.", "citations": [
            {"source_id": SECOND, "quote": "The report must include the survey findings."},
        ]}],
        "action_items": [
            {"title": "Prepare the project outline", "owner": "Alex", "due_at": None,
             "citations": [{"source_id": FIRST, "quote": QUOTE}]},
            {"title": "Review the survey findings", "owner": None,
             "due_at": "2026-10-30T18:00:00+08:00", "citations": []},
        ],
        "open_questions": ["Will the group meet on Friday?"],
        "missing_information": ["The meeting time is not confirmed."],
        "agenda": ["Review the outline", "Confirm the meeting time"],
        "follow_up": "Hello team,\nPlease review the outline and confirm a meeting time.\nThank you.",
    }


def document_text(document):
    text = [paragraph.text for paragraph in document.paragraphs]
    text.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    return "\n".join(text)


def test_word_export_contains_every_structured_field_and_only_cited_evidence():
    item, selected = summary(), sources()
    before = deepcopy((item, selected))
    output = meeting_minutes_docx(item, selected)
    assert isinstance(output, bytes) and output.startswith(b"PK")
    document = Document(BytesIO(output))
    text = document_text(document)
    assert document.paragraphs[0].style.name == "Title"
    assert document.paragraphs[0].text == item["title"]
    for expected in (
        "Summary", "Decisions", "Action items for review", "Open questions",
        "Details still to confirm", "Suggested agenda", "Follow-up draft", "Supporting text",
        item["decisions"][0]["text"], *item["open_questions"], *item["missing_information"],
        *item["agenda"], item["follow_up"], QUOTE, "Message 1", "Source 1", "Alex",
        "2026-10-26 09:00:00 SGT", "Project brief.pdf",
    ):
        assert expected in text
    assert len(document.tables) == 1
    table = document.tables[0]
    assert [cell.text for cell in table.rows[0].cells] == ["Action", "Owner", "Deadline"]
    assert table.cell(1, 1).text == "Alex"
    assert table.cell(1, 2).text == "Not specified"
    assert table.cell(2, 1).text == "Not specified"
    assert table.cell(2, 2).text == "2026-10-30\n18:00:00+08:00"
    assert table.cell(2, 2).text.replace("\n", "T") == item["action_items"][1]["due_at"]
    assert text.count(QUOTE) == 1  # Repeated citations share one full evidence excerpt.
    for excluded in (FIRST, SECOND, UNUSED, "private-account@example.invalid", "users/private-123",
                     "UNCITED PRIVATE SOURCE BODY", "PRIVATE ATTACHMENT TAIL",
                     "UNCITED SOURCE TITLE", "UNUSED SOURCE BODY"):
        assert excluded not in text
    assert (item, selected) == before


@pytest.mark.parametrize("incomplete", [False, True])
def test_partial_status_is_explicit_and_empty_fields_stay_unknown(incomplete):
    document = Document(BytesIO(meeting_minutes_docx({}, {}, incomplete=incomplete)))
    text = document_text(document)
    assert document.paragraphs[0].text == "Discussion minutes"
    assert "Not specified" in text
    assert "None recorded." in text
    assert "No supporting quotations recorded." in text
    assert ("Incomplete analysis:" in text) is incomplete
    assert not document.tables


def test_literal_markup_urls_and_commands_create_no_active_docx_content():
    literal = (
        '<script>alert("x")</script> <w:fldSimple w:instr="DDEAUTO"/> '
        "[Open](https://example.invalid/report) ![Image](https://example.invalid/image.png) "
        "https://example.invalid/path <b>literal bold markup</b>"
    )
    item = summary() | {"summary": literal, "follow_up": literal}
    output = meeting_minutes_docx(item, sources())
    document = Document(BytesIO(output))
    assert literal in document_text(document)
    with ZipFile(BytesIO(output)) as archive:
        assert not any("vba" in name.lower() or "embeddings/" in name
                       or name.startswith("word/media/") for name in archive.namelist())
        for name in archive.namelist():
            if name.endswith(".rels"):
                root = etree.fromstring(archive.read(name))
                assert not root.xpath("//*[local-name()='Relationship'][@TargetMode='External']")
        root = etree.fromstring(archive.read("word/document.xml"))
        for element in ("hyperlink", "fldSimple", "instrText", "altChunk", "object", "drawing"):
            assert not root.xpath(f"//*[local-name()='{element}']")


def test_long_multilingual_text_is_complete_and_illegal_xml_characters_are_sanitized():
    long_text = ("中文会议记录 — café résumé — العربية — 🌿\nSecond line\twith a tab. " * 150)
    illegal = "\x00\x07\x0b\x1f\ud800\ufffe\uffff"
    item = summary() | {"title": "讨论记录" + illegal, "summary": long_text + illegal,
                        "follow_up": "Follow-up end" + illegal}
    item["action_items"][0]["title"] = long_text + "LONG ACTION END"
    item["citations"] = [{"source_id": FIRST, "quote": "Cited text" + illegal + "QUOTE END"}]
    output = meeting_minutes_docx(item, sources())
    document = Document(BytesIO(output))
    text = document_text(document)
    assert long_text in text
    assert "LONG ACTION END" in document.tables[0].cell(1, 0).text
    assert "QUOTE END" in text and "Follow-up end" in text
    assert "\ufffd" in text
    assert all(character not in text for character in illegal)
    with ZipFile(BytesIO(output)) as archive:
        for name in archive.namelist():
            if name.endswith(".xml"):
                etree.fromstring(archive.read(name))


def test_a4_black_headings_and_expandable_table_layout_are_explicit():
    document = Document(BytesIO(meeting_minutes_docx(summary(), sources())))
    section = document.sections[0]
    assert abs(section.page_width - Cm(21)) < 1000
    assert abs(section.page_height - Cm(29.7)) < 1000
    assert section.left_margin == section.right_margin
    assert abs(section.left_margin - Cm(2)) < 1000
    assert document.styles["Normal"].font.size == Pt(11)
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2"):
        style = document.styles[name]
        assert str(style.font.color.rgb) == "000000"
        assert not style.element.xpath("./w:pPr/w:pBdr")
    table = document.tables[0]
    assert table.autofit is False
    widths = [column.width for column in table.columns]
    assert widths[0] > widths[2] > widths[1]
    assert abs(sum(widths) - (section.page_width - section.left_margin - section.right_margin)) < 1000
    assert table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
    assert not table._tbl.xpath(".//w:trHeight | .//w:cantSplit | .//w:noWrap")
    assert table._tbl.xpath("./w:tblPr/w:tblCellMar/w:top")
    for border in table._tbl.xpath("./w:tblPr/w:tblBorders/*"):
        assert border.get(qn("w:color")) == "D9D9D9"
    for cell in table.rows[0].cells:
        assert cell._tc.xpath("./w:tcPr/w:shd")[0].get(qn("w:fill")) == "EAF2F8"


def test_template_author_metadata_is_removed_and_export_does_not_write_files(monkeypatch):
    def reject_write(*_args, **_kwargs):
        raise AssertionError("Word export must return bytes without writing files")

    monkeypatch.setattr(Path, "write_bytes", reject_write)
    monkeypatch.setattr(Path, "write_text", reject_write)
    output = meeting_minutes_docx(summary(), sources())
    document = Document(BytesIO(output))
    assert not document.core_properties.author
    assert not document.core_properties.last_modified_by
    assert not document.core_properties.comments
    assert document.core_properties.title == summary()["title"]
    with ZipFile(BytesIO(output)) as archive:
        root = etree.fromstring(archive.read("docProps/core.xml"))
        assert not root.xpath("//*[local-name()='created' or local-name()='modified']")


def test_valid_500_character_multilingual_title_only_shortens_core_metadata():
    title = ("校园研究讨论纪要 " * 60)[:500]
    assert len(title) == 500
    document = Document(BytesIO(meeting_minutes_docx(summary() | {"title": title}, sources())))
    assert document.paragraphs[0].text == title
    assert document.core_properties.title == title[:255]


@pytest.mark.parametrize("deadline", [
    "2026-10-27T18:00:00+08:00", "2026-10-27T18:00Z", "2026-10-27T18:00:00.123-05:30",
    "2026-10-27T18:00:00", "2026-02-30T18:00:00+08:00", "Friday at 6pm TBC",
])
def test_deadline_wrapping_preserves_valid_iso_information_and_arbitrary_text(deadline):
    item = summary()
    item["action_items"][0]["due_at"] = deadline
    document = Document(BytesIO(meeting_minutes_docx(item, sources())))
    displayed = document.tables[0].cell(1, 2).text
    if deadline in {"2026-02-30T18:00:00+08:00", "Friday at 6pm TBC"}:
        assert displayed == deadline
    else:
        assert displayed.count("\n") == 1
        assert displayed.replace("\n", "T") == deadline


@pytest.mark.parametrize("agenda", [
    ["Review the outline", "Confirm the meeting time", "Agree next steps"],
    [f"Agenda item {index}" for index in range(8)],
    ["Long agenda discussion " * 40, "Another long agenda discussion " * 40],
])
def test_only_short_numbered_lists_stay_together_and_evidence_labels_keep_context(agenda):
    item = summary() | {"agenda": agenda}
    document = Document(BytesIO(meeting_minutes_docx(item, sources())))
    paragraphs = document.paragraphs
    start = next(index for index, paragraph in enumerate(paragraphs)
                 if paragraph.text == "Suggested agenda")
    entries = paragraphs[start + 1:start + 1 + len(agenda)]
    short = len(agenda) <= 5 and sum(map(len, agenda)) <= 600
    assert [paragraph.paragraph_format.keep_with_next for paragraph in entries] == (
        [True] * (len(agenda) - 1) + [False] if short else [False] * len(agenda)
    )
    for index, paragraph in enumerate(paragraphs):
        if paragraph.style.name == "Heading 2":
            assert paragraph.runs[0].bold is None  # Inherit the bold heading style.
            assert document.styles["Heading 2"].font.bold is True
            assert paragraphs[index + 1].paragraph_format.keep_with_next is True
        if paragraph.text.startswith("Quote "):
            assert paragraph.paragraph_format.keep_with_next is True
