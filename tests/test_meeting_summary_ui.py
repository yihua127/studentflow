"""Discussion outcome cards preserve unanswered questions without external writes."""

import json
from io import BytesIO

import pytest
import streamlit as st
from docx import Document
from streamlit.testing.v1 import AppTest

SCRIPT = """
import os
import json
from studentflow.assistant_ui import _render_artifacts
from studentflow.agent_tools import AssistantTools
source = {
    'source_id': 'chat-source', 'provider': 'google_chat',
    'title': 'Project discussion', 'sender_name': 'Alex',
    'sent_at': '2026-10-26T01:00:00Z',
    'text': 'Alex will clean the dataset. Check the consent form; no owner or deadline is assigned.',
}
box = AssistantTools({'sources': [source]})
outcome = json.loads(os.environ['SF_MEETING_CARD_INPUT'])
result = box.invoke(outcome['tool'], outcome['arguments'], role='meeting')
assert result['ok'], result
_render_artifacts(
    box.artifacts, {'chat-source': source}, 'meeting-test',
    analysis_status=os.environ.get('SF_MEETING_CARD_STATUS', 'completed'),
)
"""


def test_discussion_card_preserves_unknowns_and_displays_questions(monkeypatch):
    citation = {'source_id': 'chat-source', 'quote': 'Alex will clean the dataset.'}
    unassigned = {
        'source_id': 'chat-source',
        'quote': 'Check the consent form; no owner or deadline is assigned.',
    }
    monkeypatch.setenv('SF_MEETING_CARD_INPUT', json.dumps({
        'tool': 'record_meeting_summary',
        'arguments': {
            'title': 'Project discussion', 'summary': 'Preparation work and open details.',
            'citations': [citation, unassigned],
            'action_items': [
                {'title': 'Clean the dataset', 'owner': 'Alex', 'citations': [citation]},
                {'title': 'Check the consent form', 'citations': [unassigned]},
            ],
            'open_questions': ['Who will check the consent form?'],
            'missing_information': ['The consent check has no agreed deadline.'],
            'agenda': ['Confirm the remaining responsibilities.'],
            'follow_up': 'Please confirm who can check the consent form.',
        },
    }))
    app = AppTest.from_string(SCRIPT).run(timeout=30)
    assert not app.exception
    table = app.dataframe[0].value
    assert table.iloc[0]['Owner'] == 'Alex'
    assert table.iloc[1]['Owner'] == table.iloc[1]['Deadline'] == 'Not specified'
    text = '\n'.join(item.value for item in app.text)
    assert 'Who will check the consent form?' in text
    assert 'The consent check has no agreed deadline.' in text
    assert 'Please confirm who can check the consent form.' in text
    assert 'chat-source' not in text
    assert '2026-10-26 09:00:00 SGT' in text
    assert not app.button
    assert [button.label for button in app.get('download_button')] == [
        'Download minutes (Word)', 'Download minutes (TXT)', 'Download minutes (Markdown)',
    ]
    assert all(button.proto.ignore_rerun for button in app.get('download_button'))


def test_clarification_card_shows_questions_without_a_completed_summary(monkeypatch):
    monkeypatch.setenv('SF_MEETING_CARD_INPUT', json.dumps({
        'tool': 'request_meeting_clarification',
        'arguments': {'source_ids': ['chat-source'],
                      'questions': ['Which discussion should the final summary cover?']},
    }))
    app = AppTest.from_string(SCRIPT).run(timeout=30)
    assert not app.exception
    assert any('Discussion details needed' in item.value for item in app.markdown)
    assert not any('Discussion summary' in item.value for item in app.markdown)
    assert 'Which discussion should the final summary cover?' in [item.value for item in app.text]
    assert not app.dataframe
    assert not app.button
    assert not app.get('download_button')


@pytest.mark.parametrize('status', ['completed', 'partial'])
def test_downloads_contain_current_minutes_and_preserve_run_status(monkeypatch, status):
    downloads = []
    original = st.download_button

    def capture(label, data, **kwargs):
        downloads.append({'label': label, 'data': data, **kwargs})
        return original(label, data, **kwargs)

    monkeypatch.setattr(st, 'download_button', capture)
    monkeypatch.setenv('SF_MEETING_CARD_STATUS', status)
    monkeypatch.setenv('SF_MEETING_CARD_INPUT', json.dumps({
        'tool': 'record_meeting_summary',
        'arguments': {
            'title': 'Current discussion', 'summary': 'Check the consent form.',
            'citations': [{'source_id': 'chat-source',
                          'quote': 'Check the consent form; no owner or deadline is assigned.'}],
            'action_items': [{'title': 'Check the consent form', 'citations': [{
                'source_id': 'chat-source',
                'quote': 'Check the consent form; no owner or deadline is assigned.',
            }]}],
            'open_questions': ['Who will check the form?'],
            'follow_up': 'Please confirm the owner.',
        },
    }))
    app = AppTest.from_string(SCRIPT).run(timeout=30)
    assert not app.exception
    word, plain, markdown = downloads
    assert word['file_name'].endswith('.docx')
    assert word['mime'] == (
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    assert word['type'] == 'primary'
    document = Document(BytesIO(word['data']))
    text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
    assert 'Current discussion' in text and 'Who will check the form?' in text
    assert 'Please confirm the owner.' in text
    assert document.tables[0].rows[1].cells[1].text == 'Not specified'
    assert document.tables[0].rows[1].cells[2].text == 'Not specified'
    assert 'chat-source' not in text
    assert 'Message 1' in text and '2026-10-26 09:00:00 SGT' in text
    assert plain['file_name'].endswith('.txt')
    assert markdown['file_name'].endswith('.md')
    for download in downloads:
        assert download['on_click'] == 'ignore'
    for content in (text, plain['data'], markdown['data']):
        assert ('incomplete' in content.casefold()) is (status == 'partial')


def test_multiple_minutes_have_independent_download_buttons():
    script = """
from studentflow.assistant_ui import _render_artifacts
_render_artifacts({'meeting_summaries': [
    {'id': 'first', 'title': 'First discussion', 'summary': 'First outcome'},
    {'id': 'second', 'title': 'Second discussion', 'summary': 'Second outcome'},
]}, {}, 'report')
"""
    app = AppTest.from_string(script).run(timeout=30)
    assert not app.exception
    downloads = app.get('download_button')
    assert len(downloads) == 6
    assert len({button.key for button in downloads}) == 6
