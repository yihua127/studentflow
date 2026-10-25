"""Editable Word minutes built entirely in memory from the reviewed result.

Model and source strings are inserted as literal text, never HTML, links, fields
or external media. The common export model controls evidence scope and labels.
"""

from __future__ import annotations

import re
from datetime import datetime
from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from studentflow.meeting_export import meeting_minutes_document

_ILLEGAL_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_COLUMN_WIDTHS = (Cm(9.2), Cm(3.0), Cm(4.8))
_ISO_DEADLINE = re.compile(
    r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}(?::\d{2}(?:\.\d+)?)?(?:Z|[+-]\d{2}:\d{2})?"
)


def _text(value: str) -> str:
    # Preserve valid Unicode, line breaks and tabs while avoiding invalid XML.
    return _ILLEGAL_XML.sub("\ufffd", value)


def _paragraph(container, text: str, *, style=None, bold=None):
    paragraph = container.add_paragraph(style=style)
    paragraph.add_run(_text(text)).bold = bold
    return paragraph


def _deadline_text(value: str) -> str:
    """Wrap a valid ISO timestamp deliberately without changing its information."""
    if _ISO_DEADLINE.fullmatch(value):
        try:
            datetime.fromisoformat(value)
        except ValueError:
            pass
        else:
            return value.replace("T", "\n", 1)
    return value


def _numbered_list(document, items: list[str]) -> None:
    keep_group = len(items) <= 5 and sum(len(item) for item in items) <= 600
    for index, text in enumerate(items, 1):
        paragraph = _paragraph(document, f"{index}. {text}")
        paragraph.paragraph_format.keep_with_next = keep_group and index < len(items)
        paragraph.paragraph_format.keep_together = keep_group
    if not items:
        _paragraph(document, "None recorded.")


def _references(container, references: list[str]) -> None:
    if references:
        paragraph = _paragraph(container, "Supporting sources: " + ", ".join(references))
        for run in paragraph.runs:
            run.italic = True


def _document_styles(document) -> None:
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        fonts.set(qn("w:cs"), "Arial")
        paragraph = style.paragraph_format
        paragraph.line_spacing = 1.12
        paragraph.space_after = Pt(6)
        paragraph.widow_control = True
    normal = document.styles["Normal"].paragraph_format
    normal.keep_together = False
    normal.keep_with_next = False
    title = document.styles["Title"]
    title.font.size = Pt(22)
    title.paragraph_format.space_after = Pt(6)
    title.paragraph_format.keep_with_next = True
    document.styles["Subtitle"].paragraph_format.space_after = Pt(12)
    for name, size in (("Heading 1", 13), ("Heading 2", 11)):
        style = document.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.keep_with_next = True
    # Remove template-specific title rules/theme colors instead of inheriting
    # formatting that can appear differently when the file is opened in Word.
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2"):
        style = document.styles[name]
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
    core = document.core_properties
    core.author = ""
    core.last_modified_by = ""
    core.comments = ""
    core.subject = ""
    core.keywords = ""
    for timestamp in core._element.xpath("./dcterms:created | ./dcterms:modified | ./cp:lastPrinted"):
        timestamp.getparent().remove(timestamp)


def _table_properties(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    props = table._tbl.tblPr
    width = props.find(qn("w:tblW"))
    width.set(qn("w:w"), str(Cm(17).twips))
    width.set(qn("w:type"), "dxa")
    borders = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + name)
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            border.set(qn("w:" + key), value)
        borders.append(border)
    props.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for name, value in (("top", "100"), ("bottom", "100"), ("left", "120"), ("right", "120")):
        margin = OxmlElement("w:" + name)
        margin.set(qn("w:w"), value)
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    props.append(margins)
    for column, width in zip(table.columns, _COLUMN_WIDTHS, strict=True):
        column.width = width


def _action_table(document, actions: list[dict]) -> None:
    table = document.add_table(rows=1, cols=3)
    _table_properties(table)
    header = table.rows[0]
    repeating = OxmlElement("w:tblHeader")
    repeating.set(qn("w:val"), "true")
    header._tr.get_or_add_trPr().append(repeating)
    for cell, label in zip(header.cells, ("Action", "Owner", "Deadline"), strict=True):
        cell.text = label
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "EAF2F8")
        cell._tc.get_or_add_tcPr().append(shading)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for action in actions:
        cells = table.add_row().cells
        values = (action["title"], action["owner"], _deadline_text(action["deadline"]))
        for cell, value in zip(cells, values, strict=True):
            cell.text = _text(value)
        _references(cells[0], action["references"])
    for row in table.rows:
        # No fixed heights or cantSplit: large/wrapped cells may flow to another
        # page with a repeated header instead of clipping or creating a gap.
        for cell, width in zip(row.cells, _COLUMN_WIDTHS, strict=True):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = False
                paragraph.paragraph_format.keep_together = False
                paragraph.paragraph_format.space_after = Pt(3)
        for cell in row.cells[1:]:
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def meeting_minutes_docx(item: dict, sources: dict, *, incomplete: bool = False) -> bytes:
    """Return a self-contained editable DOCX without opening or writing files."""
    content = meeting_minutes_document(item, sources, incomplete=incomplete)
    document = Document()
    _document_styles(document)
    # python-docx limits core-property strings to 255 characters. The visible
    # title remains complete, including valid 500-character model titles.
    document.core_properties.title = _text(content["title"])[:255]
    _paragraph(document, content["title"], style="Title")
    _paragraph(document, "Discussion minutes", style="Subtitle")
    if content["warning"]:
        _paragraph(document, content["warning"], bold=True)

    document.add_heading("Summary", level=1)
    _paragraph(document, content["summary"])
    _references(document, content["summary_references"])

    document.add_heading("Decisions", level=1)
    for index, decision in enumerate(content["decisions"], 1):
        _paragraph(document, f"{index}. {decision['text']}")
        _references(document, decision["references"])
    if not content["decisions"]:
        _paragraph(document, "None recorded.")

    document.add_heading("Action items for review", level=1)
    if content["action_items"]:
        _action_table(document, content["action_items"])
    else:
        _paragraph(document, "None recorded.")

    for key, heading in (
        ("open_questions", "Open questions"),
        ("missing_information", "Details still to confirm"),
        ("agenda", "Suggested agenda"),
    ):
        document.add_heading(heading, level=1)
        _numbered_list(document, content[key])

    document.add_heading("Follow-up draft", level=1)
    _paragraph(document, content["follow_up"])
    document.add_heading("Supporting text", level=1)
    for source in content["supporting_text"]:
        _paragraph(document, source["label"], style="Heading 2")
        metadata = _paragraph(document, source["heading"])
        metadata.paragraph_format.keep_with_next = True
        for index, quote in enumerate(source["quotes"], 1):
            label = _paragraph(document, f"Quote {index}", bold=True)
            label.paragraph_format.keep_with_next = True
            paragraph = _paragraph(document, quote)
            paragraph.paragraph_format.left_indent = Cm(0.5)
    if not content["supporting_text"]:
        _paragraph(document, "No supporting quotations recorded.")

    output = BytesIO()
    document.save(output)
    return output.getvalue()
